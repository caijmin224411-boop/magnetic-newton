import os
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "01_theory"
ASSET_DIR = ROOT / "07_generated_assets" / "formula_rendered"
DOCX_PATH = OUT_DIR / "Magnetic_Newtons_Cradle_Theory_Formulas.docx"

PY_RUNTIME = Path(os.environ.get(
    "CODEX_PYTHON",
    r"C:\Users\66\.cache\codex-runtimes\codex-primary-runtime\dependencies\python\python.exe",
))
NODE_RUNTIME = Path(os.environ.get(
    "CODEX_NODE",
    r"C:\Users\66\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\bin\node.exe",
))
NODE_MODULES = Path(os.environ.get(
    "CODEX_NODE_MODULES",
    r"C:\Users\66\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\node_modules",
))


FORMULAS = {
    "f01_pivot": r"X_i=\left(i-\frac{N+1}{2}\right)d",
    "f02_position": r"x_i=X_i+L_i\sin\theta_i,\qquad y_i=-L_i\cos\theta_i",
    "f03_state": r"\mathbf{s}=(\theta_1,\ldots,\theta_5,\omega_1,\ldots,\omega_5)^T",
    "f04_gravity": r"\boldsymbol{G}_i=(0,-mg),\qquad \tau_{g,i}=-mgL_i\sin\theta_i",
    "f05_damping": r"\tau_{d,i}=-c_i\omega_i",
    "f06_relative": r"\Delta x_{ij}=x_i-x_j,\qquad \Delta y_{ij}=y_i-y_j",
    "f07_gap": r"g_{ij}=|\Delta x_{ij}|-2R,\qquad o_{ij}=|\Delta y_{ij}|",
    "f08_force_table": r"F_{\mathrm{axial}}=F_{\mathrm{axial}}(g_{ij},o_{ij}),\qquad F_{\mathrm{lateral}}=F_{\mathrm{lateral}}(g_{ij},o_{ij})",
    "f09_force_components": r"F_{x,ij}=\operatorname{sgn}(\Delta x_{ij})F_{\mathrm{axial}},\qquad F_{y,ij}=\operatorname{sgn}(\Delta y_{ij})F_{\mathrm{lateral}}",
    "f10_dipole_scale": r"F_{\mathrm{mag}}\propto\frac{\mu_1\mu_2}{r^4},\qquad F_{\mathrm{mag}}\propto B^2",
    "f11_arm": r"\boldsymbol{r}_i=(L_i\sin\theta_i,-L_i\cos\theta_i)",
    "f12_magnetic_torque": r"\tau_{m,ij}=(\boldsymbol{r}_i\times\boldsymbol{F}_{ij})_z=r_xF_{y,ij}-r_yF_{x,ij}",
    "f13_magnetic_torque_expanded": r"\tau_{m,ij}=L_i\sin\theta_iF_{y,ij}+L_i\cos\theta_iF_{x,ij}",
    "f14_total_torque": r"\tau_i=-mgL_i\sin\theta_i-c_i\omega_i+\sum_j\left(L_i\sin\theta_iF_{y,ij}+L_i\cos\theta_iF_{x,ij}\right)",
    "f15_inertia": r"I_i=mL_i^2",
    "f16_motion": r"\theta_i''=-\frac{g}{L_i}\sin\theta_i-\frac{c_i}{mL_i^2}\theta_i'+\frac{1}{mL_i}\sum_j\left(\sin\theta_iF_{y,ij}+\cos\theta_iF_{x,ij}\right)",
    "f17_state_equations": r"\dot{\theta}_i=\omega_i,\qquad \dot{\omega}_i=-\frac{g}{L_i}\sin\theta_i-\frac{c_i}{mL_i^2}\omega_i+\frac{1}{mL_i}\sum_j\left(\sin\theta_iF_{y,ij}+\cos\theta_iF_{x,ij}\right)",
    "f18_equilibrium": r"-mgL_i\sin\theta_{i,\mathrm{eq}}+\sum_j\left(L_i\sin\theta_{i,\mathrm{eq}}F_{y,ij}+L_i\cos\theta_{i,\mathrm{eq}}F_{x,ij}\right)=0",
    "f19_linearization": r"\theta_i(t)=\theta_{i,\mathrm{eq}}+q_i(t),\qquad |q_i|\ll1",
    "f20_matrix": r"M\ddot{\boldsymbol{q}}+C\dot{\boldsymbol{q}}+K\boldsymbol{q}=0",
    "f21_stiffness": r"K_{ij}=-\left.\frac{\partial\tau_i}{\partial\theta_j}\right|_{\boldsymbol{\theta}=\boldsymbol{\theta}_{\mathrm{eq}}}",
    "f22_eigen": r"K\boldsymbol{\phi}_k=\Omega_k^2M\boldsymbol{\phi}_k,\qquad f_k=\frac{\Omega_k}{2\pi}",
    "f23_damped": r"q_k(t)=A_ke^{-\gamma_kt}\cos(\Omega_{d,k}t+\varphi_k),\qquad \Omega_{d,k}=\sqrt{\Omega_k^2-\gamma_k^2}",
    "f24_fft": r"A_i(f)=\left|\mathcal{F}\{\theta_i(t)-\overline{\theta_i}\}\right|",
    "f25_param": r"f_k\approx\frac{1}{2\pi}\sqrt{\lambda_k},\qquad \lambda_k=\lambda_k(g,L_i,m,B,d,R,c_i,\theta_0)",
}


SECTIONS = [
    ("基础定义", [
        ("悬点位置", "f01_pivot", "d 是相邻悬点间距，N=5。"),
        ("磁体质心坐标", "f02_position", "L_i 是第 i 个磁体的有效摆长。"),
        ("系统状态变量", "f03_state", "五摆系统共有 10 维相空间。"),
    ]),
    ("基础受力", [
        ("重力与重力力矩", "f04_gravity", "悬线拉力通过悬点，不产生绕悬点的力矩。"),
        ("阻尼力矩", "f05_damping", "c_i 可由单摆自由衰减实验拟合。"),
    ]),
    ("磁力大小与方向", [
        ("相对位置", "f06_relative", "磁力由两个磁体的瞬时相对位置决定。"),
        ("间隙与错位", "f07_gap", "g_ij 是表面间隙，o_ij 是竖直错位。"),
        ("有限圆柱磁体力表", "f08_force_table", "近距离时不用点偶极近似，而使用有限圆柱磁体力表插值。"),
        ("磁力分量", "f09_force_components", "符号由相对位置决定，大小由力表读取。"),
        ("点偶极量级关系", "f10_dipole_scale", "该式只用于判断趋势：磁力近似随 B^2 增强。"),
    ]),
    ("磁力力矩与总力矩", [
        ("力臂", "f11_arm", "力臂从悬点指向磁体质心。"),
        ("磁力力矩", "f12_magnetic_torque", "二维运动中只需要 z 方向力矩。"),
        ("展开形式", "f13_magnetic_torque_expanded", "F_x 和 F_y 都会贡献角加速度。"),
        ("总力矩", "f14_total_torque", "求和 j 当前取相邻磁体 i-1 和 i+1。"),
    ]),
    ("总运动方程", [
        ("转动惯量", "f15_inertia", "把磁体近似为集中质量。"),
        ("二阶非线性运动方程", "f16_motion", "这是当前数值模型的核心方程。"),
        ("一阶状态方程", "f17_state_equations", "数值积分时使用状态空间形式。"),
    ]),
    ("静平衡与线性化", [
        ("静平衡方程", "f18_equilibrium", "静止时角速度和角加速度都为零。"),
        ("小振动展开", "f19_linearization", "在静平衡附近令 theta 等于平衡角加小扰动。"),
        ("矩阵方程", "f20_matrix", "M、C、K 分别是质量矩阵、阻尼矩阵和刚度矩阵。"),
        ("刚度矩阵", "f21_stiffness", "磁力对角度的导数进入耦合刚度。"),
    ]),
    ("运动方程解与解析", [
        ("无阻尼模态解", "f22_eigen", "频率来自广义本征值问题。"),
        ("有阻尼近似解", "f23_damped", "阻尼小时主要改变峰宽和衰减速度。"),
        ("频谱提取", "f24_fft", "模型和 Tracker 都通过 FFT 提取主峰。"),
        ("参数关系", "f25_param", "lambda_k 同时包含重力刚度和磁耦合刚度。"),
    ]),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(10)
    if color:
        r.font.color.rgb = RGBColor(*color)


def style_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.size = Pt(22)
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 1"].font.color.rgb = RGBColor(28, 63, 82)
    styles["Heading 2"].font.color.rgb = RGBColor(70, 70, 70)

    header = section.header.paragraphs[0]
    header.text = "Magnetic Newton's Cradle - Theory"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(120, 120, 120)


def render_formulas():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    tex_dir = ASSET_DIR / "tex"
    tex_dir.mkdir(parents=True, exist_ok=True)
    rendered = {}
    for key, formula in FORMULAS.items():
        tex_path = tex_dir / f"{key}.tex"
        dvi_path = tex_dir / f"{key}.dvi"
        svg_path = ASSET_DIR / f"{key}.svg"
        png_path = ASSET_DIR / f"{key}.png"
        tex = "\n".join([
            r"\documentclass{article}",
            r"\usepackage{amsmath,amssymb}",
            r"\pagestyle{empty}",
            r"\begin{document}",
            rf"\begin{{displaymath}}{formula}\end{{displaymath}}",
            r"\end{document}",
        ])
        tex_path.write_text(tex, encoding="ascii")
        subprocess.run([
            "latex", "-interaction=nonstopmode", "-halt-on-error",
            f"-output-directory={tex_dir}", str(tex_path),
        ], check=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
        subprocess.run([
            "dvisvgm", "--bbox=min", "--exact", "--no-fonts",
            "-o", str(svg_path), str(dvi_path),
        ], check=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
        js = (
            "const sharp=require('sharp');"
            f"sharp({str(svg_path).replace(chr(92), '/').__repr__()})"
            ".resize({height:130,withoutEnlargement:false})"
            f".png().toFile({str(png_path).replace(chr(92), '/').__repr__()})"
            ".then(()=>{}).catch(e=>{console.error(e);process.exit(1);});"
        )
        env = os.environ.copy()
        env["NODE_PATH"] = str(NODE_MODULES)
        subprocess.run([str(NODE_RUNTIME), "-e", js], check=True, env=env)
        rendered[key] = png_path
    return rendered


def add_formula(doc, image_path, note=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(13.8))
    if note:
        p = doc.add_paragraph(note)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(100, 100, 100)


def build_docx(rendered):
    doc = Document()
    style_doc(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("磁力牛顿摆理论推导").bold = True
    sub = doc.add_paragraph("基础定义、受力分析、磁力分解、总运动方程与模态解析")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in sub.runs:
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(90, 90, 90)

    doc.add_paragraph(
        "本文档对应当前五磁摆模型：每个磁体被视为一个平面单摆，磁力由有限圆柱磁体力表给出，"
        "通过力矩方程建立非线性运动方程。公式均由 LaTeX 编译为图片后嵌入，避免显示源代码。"
    )

    meta = doc.add_table(rows=1, cols=4)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["对象", "模型", "自由度", "核心输出"]
    vals = ["五磁体牛顿摆", "非线性耦合摆阵列", "10 维相空间", "静平衡、频谱、模态"]
    for i, h in enumerate(headers):
        cell = meta.cell(0, i)
        set_cell_shading(cell, "EAF2F4")
        set_cell_text(cell, f"{h}\n{vals[i]}", bold=True, color=(30, 65, 80))
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for sec_title, items in SECTIONS:
        doc.add_heading(sec_title, level=1)
        for title, key, note in items:
            doc.add_heading(title, level=2)
            add_formula(doc, rendered[key], note)

    doc.add_heading("参数对频率的解析关系", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(["参数增大", "主要影响", "频率趋势"]):
        set_cell_shading(hdr[i], "DDECEF")
        set_cell_text(hdr[i], text, bold=True, color=(25, 65, 82))
    rows = [
        ("有效摆长 L_i", "重力项 g/L_i 与磁力项 1/(mL_i)", "低频主峰下降"),
        ("悬点间距 d", "改变磁体间隙和磁力梯度", "间距变大时磁耦合峰下降"),
        ("磁场 B", "磁力近似随 B^2 增强", "磁耦合峰升高"),
        ("质量 m", "磁力角加速度与 1/m 成正比", "磁耦合影响减弱"),
        ("阻尼 c_i", "能量耗散和峰宽", "主频小变，峰变宽"),
        ("释放角 theta_0", "非线性激发强度", "大角度时频谱更复杂"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text)

    doc.add_heading("结论", level=1)
    p = doc.add_paragraph()
    p.add_run("核心公式：").bold = True
    p.add_run(" 磁力牛顿摆的频率来自线性化后的广义本征值问题，而不是单个普通单摆公式。")
    add_formula(doc, rendered["f22_eigen"], "重力刚度和磁力耦合刚度共同决定各个模态频率。")

    doc.save(DOCX_PATH)
    return DOCX_PATH


def main():
    rendered = render_formulas()
    out = build_docx(rendered)
    print(out)


if __name__ == "__main__":
    main()
